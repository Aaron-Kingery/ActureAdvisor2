import io
import re
from typing import Optional
import fitz  # pymupdf
import docx
import openpyxl

class DocumentProcessor:
    @staticmethod
    def extract_text(file_stream: io.BytesIO, file_type: str) -> str:
        """
        Extracts text from various file formats.
        """
        file_type = file_type.lower()

        text = ""
        if 'pdf' in file_type:
            text = DocumentProcessor._extract_from_pdf(file_stream)
        elif 'docx' in file_type or 'word' in file_type:
            text = DocumentProcessor._extract_from_docx(file_stream)
        elif 'xlsx' in file_type or 'excel' in file_type or 'spreadsheet' in file_type:
            text = DocumentProcessor._extract_from_xlsx(file_stream)
        elif 'text' in file_type or 'md' in file_type or 'txt' in file_type:
            text = file_stream.read().decode('utf-8', errors='ignore')
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

        # Sanitize: Remove null bytes which cause PostgreSQL errors
        if text:
            text = text.replace('\x00', '')

        return text

    @staticmethod
    def _extract_from_pdf(stream: io.BytesIO) -> str:
        """
        Structure-aware PDF extraction.
        Uses font-size heuristics to detect headings and preserve document structure.
        """
        sections = []
        with fitz.open(stream=stream, filetype="pdf") as doc:
            for page in doc:
                blocks = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]
                for block in blocks:
                    if block["type"] != 0:  # skip image blocks
                        continue
                    for line in block["lines"]:
                        line_text = ""
                        max_font_size = 0
                        for span in line["spans"]:
                            line_text += span["text"]
                            max_font_size = max(max_font_size, span["size"])

                        line_text = line_text.strip()
                        if not line_text:
                            continue

                        # Font-size heuristics for heading detection
                        if max_font_size >= 18:
                            sections.append(f"\n# {line_text}")
                        elif max_font_size >= 14:
                            sections.append(f"\n## {line_text}")
                        elif max_font_size >= 12 and line["spans"][0].get("flags", 0) & 2 ** 4:
                            # Bold text at 12pt+ treated as H3
                            sections.append(f"\n### {line_text}")
                        else:
                            sections.append(line_text)

                sections.append("")  # page break

        return "\n".join(sections)

    @staticmethod
    def _extract_from_docx(stream: io.BytesIO) -> str:
        """
        DOCX extraction that preserves headings and captures table content.
        Iterates paragraphs and tables in document order.
        """
        doc = docx.Document(stream)
        parts = []

        # Build an ordered list of all block-level elements (paragraphs + tables)
        # by walking the document body XML in order
        from docx.oxml.ns import qn
        body = doc.element.body
        para_index = 0
        table_index = 0

        paragraphs = doc.paragraphs
        tables = doc.tables

        for child in body:
            if child.tag == qn('w:p'):
                if para_index < len(paragraphs):
                    para = paragraphs[para_index]
                    para_text = para.text.strip()
                    if para_text:
                        style_name = (para.style.name or "").lower() if para.style else ""
                        if 'heading 1' in style_name:
                            parts.append(f"\n# {para_text}")
                        elif 'heading 2' in style_name:
                            parts.append(f"\n## {para_text}")
                        elif 'heading 3' in style_name:
                            parts.append(f"\n### {para_text}")
                        elif 'heading 4' in style_name:
                            parts.append(f"\n#### {para_text}")
                        elif 'title' in style_name:
                            parts.append(f"\n# {para_text}")
                        else:
                            parts.append(para_text)
                para_index += 1

            elif child.tag == qn('w:tbl'):
                if table_index < len(tables):
                    table = tables[table_index]
                    table_md = DocumentProcessor._table_to_markdown(table)
                    if table_md:
                        parts.append(table_md)
                table_index += 1

        return "\n".join(parts)

    @staticmethod
    def _table_to_markdown(table) -> str:
        """Convert a python-docx Table object to a markdown table string."""
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = cell.text.strip().replace('\n', ' ').replace('|', '\\|')
                cells.append(cell_text)
            rows.append(cells)

        if not rows:
            return ""

        # Determine column count from the widest row
        max_cols = max(len(r) for r in rows)

        # Pad short rows
        for r in rows:
            while len(r) < max_cols:
                r.append("")

        lines = []
        # Header row
        lines.append("| " + " | ".join(rows[0]) + " |")
        # Separator
        lines.append("| " + " | ".join(["---"] * max_cols) + " |")
        # Data rows
        for row in rows[1:]:
            # Skip completely empty rows
            if all(not c for c in row):
                continue
            lines.append("| " + " | ".join(row) + " |")

        return "\n" + "\n".join(lines) + "\n"

    @staticmethod
    def _extract_from_xlsx(stream: io.BytesIO) -> str:
        wb = openpyxl.load_workbook(stream, data_only=True)
        text = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            text.append(f"Sheet: {sheet}")

            rows = list(ws.rows)
            if not rows:
                continue

            # Helper to format row
            def format_row(row_cells):
                cells = [str(cell.value).strip().replace('\n', ' ') if cell.value is not None else "" for cell in row_cells]
                return "| " + " | ".join(cells) + " |"

            # Header
            header_row = rows[0]
            text.append(format_row(header_row))

            # Separator
            separator = "| " + " | ".join(["---"] * len(header_row)) + " |"
            text.append(separator)

            # Data
            for row in rows[1:]:
                # skip empty rows
                if all(cell.value is None for cell in row):
                    continue
                text.append(format_row(row))

            text.append("\n")  # Spacing between sheets

        return "\n".join(text)
